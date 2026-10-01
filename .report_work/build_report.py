from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports"
OUT.mkdir(exist_ok=True)
FIG = OUT / "人机交互技术流程图.png"
DOC = OUT / "机器人语音人机交互流程技术验证报告.docx"

font_path = r"C:\Windows\Fonts\msyh.ttc"
font_bold_path = r"C:\Windows\Fonts\msyhbd.ttc"
f33 = ImageFont.truetype(font_bold_path, 33)
f30 = ImageFont.truetype(font_path, 30)
f27 = ImageFont.truetype(font_path, 27)
f25 = ImageFont.truetype(font_path, 25)

im = Image.new("RGB", (2000, 1770), "white")
d = ImageDraw.Draw(im)
navy = "#17324D"
blue = "#2A6F9E"
light = "#EAF3F9"
gray = "#59636D"
orange = "#A75323"

def centered(lines, xy, font=f30, fill=navy, spacing=8):
    x0, y0, x1, y1 = xy
    if isinstance(lines, str):
        lines = lines.split("\n")
    hs = [d.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(hs) + spacing * (len(lines) - 1)
    y = (y0 + y1 - total) / 2
    for line, h in zip(lines, hs):
        w = d.textlength(line, font=font)
        d.text(((x0 + x1 - w) / 2, y), line, font=font, fill=fill)
        y += h + spacing

def box(xy, lines, fill=light, stroke=blue, font=f30):
    d.rounded_rectangle(xy, radius=18, fill=fill, outline=stroke, width=4)
    centered(lines, xy, font=font)

def diamond(cx, cy, rx, ry, lines):
    pts = [(cx, cy-ry), (cx+rx, cy), (cx, cy+ry), (cx-rx, cy)]
    d.polygon(pts, fill="#FFF6E9", outline=orange, width=4)
    centered(lines, (cx-rx+35, cy-ry+20, cx+rx-35, cy+ry-20), font=f27)

def arrow(points, color=blue, width=5, head=17):
    d.line(points, fill=color, width=width, joint="curve")
    (x0, y0), (x1, y1) = points[-2:]
    import math
    a = math.atan2(y1-y0, x1-x0)
    sides = []
    for da in (2.55, -2.55):
        sides.append((x1 + head*math.cos(a+da), y1 + head*math.sin(a+da)))
    d.polygon([(x1,y1), *sides], fill=color)

center = 950
d.text((70, 30), "机器人语音人机交互技术流程", font=f33, fill=navy)
box((520, 100, 1380, 205), "Snowboy 检测唤醒词")
box((520, 255, 1380, 375), "首次录音和 Vosk 语音转写")
box((520, 425, 1380, 550), "加载关键词映射并初始化串口", font=f27)
diamond(center, 665, 270, 80, "命中交互模式？")
box((1510, 615, 1915, 715), "其他功能或\n未匹配处理", fill="#F2F4F5", stroke="#858D94", font=f25)
box((520, 805, 1380, 920), "录音并转写当前问题")
diamond(center, 1040, 300, 80, "包含“结束”\n或“退出”？")
box((1510, 985, 1915, 1095), "退出交互循环\n并关闭串口", fill="#F2F4F5", stroke="#858D94", font=f25)
box((520, 1170, 1380, 1300), "interaction 调用本地 Ollama\nqwen:0.5b，流式生成", font=f27)
box((520, 1350, 1380, 1505), "逐块拼接原文；清理播报文本\n遇标点或缓冲区超过 30 字时同步播报", font=f25)
box((520, 1555, 1380, 1690), "播报剩余文本；打印耗时和字符数", font=f27)

for ya, yb in [(205,255),(375,425),(550,585),(745,805),(920,960),(1120,1170),(1300,1350),(1505,1555)]:
    arrow([(center,ya),(center,yb)])
arrow([(1220,665),(1510,665)], color="#858D94")
d.text((1290,625), "否", font=f25, fill=gray)
d.text((985,760), "是", font=f25, fill=gray)
arrow([(1250,1040),(1510,1040)], color="#858D94")
d.text((1345,998), "是", font=f25, fill=gray)
d.text((985,1128), "否", font=f25, fill=gray)
arrow([(520,1620),(255,1620),(255,862),(520,862)])
d.text((270,1470), "下一轮提问", font=f25, fill=blue)
im.save(FIG, dpi=(300,300))

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)

def style(name, size, bold=False, before=0, after=0):
    s = doc.styles[name]
    s.font.name = "Microsoft YaHei"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    return s

style("Title", 20, True, 0, 10)
style("Normal", 10.5, False, 0, 7)
style("Heading 1", 13, True, 14, 7)
style("Heading 2", 11, True, 9, 5)
style("Caption", 9, False, 5, 6)

title = doc.add_paragraph(style="Title")
title.add_run("机器人语音人机交互流程技术验证报告")
p = doc.add_paragraph()
p.add_run("验证对象：").bold = True
p.add_run("现有机器人程序中的语音问答流程，重点检查 interaction_zh.py 与其入口调用。")
p = doc.add_paragraph()
p.add_run("结论：").bold = True
p.add_run("代码已形成“语音识别—本地模型生成—语音播报”的处理链。静态检查和隔离测试通过；当前环境未连接机器人及其本地 Ollama 服务，因此不把整机运行效果列为已验证。")

doc.add_heading("一  流程范围", level=1)
doc.add_paragraph("主程序通过 Snowboy 检测唤醒词，随后调用 Vosk 把语音转成文字。action_neo.py 根据 speech_train.json 匹配“交互”指令，进入连续提问。用户说“结束”或“退出”时离开该模式。")
doc.add_heading("二  技术流程图", level=1)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
p.add_run().add_picture(str(FIG), width=Cm(16.2))
p = doc.add_paragraph(style="Caption")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("流程图依据当前代码绘制，箭头表示实际调用顺序")

doc.add_page_break()
doc.add_heading("三  关键环节说明", level=1)
items = [
    ("入口和分流", "snowboydecoder.py 在检测到唤醒词后调用 robot_process()，先执行 speech_recog()，再交由 action_neo() 匹配关键词。action_neo() 会先读取 speech_train.json 并初始化串口，然后进入 interaction_mode 分支。"),
    ("语音转写", "speech_recog.py 使用名为 snowboy_mic 的输入设备，以 16 kHz、单声道、16 位采集音频，保存为临时 WAV 文件，再交给本地 Vosk 模型识别。代码传入的录音参数为 3；按每次读取 2000 帧、循环 30 次计算，理论采样时长约 3.75 秒，实际耗时还受设备读入影响。"),
    ("模型请求", "interaction_zh.py 创建指向 127.0.0.1:11434 的 Ollama 客户端，调用 chat 接口，模型名为 qwen:0.5b，开启流式返回。每次请求只发送本轮问题，代码没有附带历史消息。num_predict 为 128，temperature 为 0.7，num_ctx 为 512，num_thread 为 2。"),
    ("输出与播报", "收到数据块后，程序从 message.content 读取文本并拼接成完整回复。播报前只去除 * 和 #；当新块含句末标点或换行，或者缓冲文本超过 30 字时，调用 espeak 中文语音播报。播报调用为同步执行，完成后才继续读取后续数据块。流结束后还会播报剩余缓冲文本。"),
    ("退出和异常", "交互循环在识别文本包含“结束”或“退出”时退出。Ollama 请求建立失败时函数返回错误文字；流迭代中断时会打印错误，播报已有缓冲并返回已收到的内容。上层只打印耗时和 len(response)，该值是字符数，不是模型 token 数。"),
]
for head, body in items:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    r = p.add_run(head + "。")
    r.bold = True
    p.add_run(body)

doc.add_heading("四  验证记录", level=1)
rows = [
    ("检查项", "方法", "结果"),
    ("代码及词表可解析", "编译检查 4 个 Python 文件，解析 speech_train.json", "通过"),
    ("模型请求参数", "隔离 Ollama 客户端，检查地址、模型、单轮消息和流式标志", "通过"),
    ("流式处理", "输入两段模拟回复，检查原文拼接、标记清理和剩余文本播报", "通过"),
    ("流中断", "模拟已收到一段内容后抛出异常", "返回部分内容，并播报缓冲"),
    ("整机语音交互", "需在机器人上连同麦克风、串口、Ollama 和 espeak 测试", "未执行"),
]
table = doc.add_table(rows=len(rows), cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for col, width in zip(table.columns, [Cm(4.0), Cm(9.0), Cm(3.5)]):
    col.width = width
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = table.cell(i,j)
        cell.text = val
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "DCEAF2" if i == 0 else ("F7FAFC" if i%2 == 0 else "FFFFFF"))
        tcPr.append(shd)
        mar = OxmlElement("w:tcMar")
        for side in ("top","left","bottom","right"):
            m = OxmlElement("w:"+side)
            m.set(qn("w:w"), "90")
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tcPr.append(mar)
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(0)
            for run in para.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(9)
                run.bold = (i == 0)
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        borders = OxmlElement("w:tcBorders")
        for side in ("top","left","bottom","right"):
            b = OxmlElement("w:"+side)
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), "D9D9D9")
            borders.append(b)
        tcPr.append(borders)

doc.add_heading("五  结论和待验证点", level=1)
doc.add_paragraph("现有代码具备本地语音问答的基本链路，隔离测试覆盖了 Ollama 调用格式和流式回复处理。交付前仍需在目标机器人上完成端到端测试，记录唤醒、首句播报、连续提问、退出和再次唤醒的实际结果及耗时。")
doc.add_paragraph("当前实现还存在三处需要重点核对：一是交互模式进入前必须成功打开串口；二是 Snowboy 的输入流在唤醒后被关闭，代码未见重新打开，退出交互后能否再次唤醒需现场确认；三是机器人入口没有加入命令行测试中的“红莲”角色说明，模型回答不保证带有该身份设定。")

doc.save(DOC)
print(DOC)
print(FIG)
